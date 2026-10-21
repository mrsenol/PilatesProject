#!/usr/bin/env python3
"""Create a self-contained archive of canonical Reformer Pilates solution runs.

The script deliberately excludes smoke tests, partial diagnostics and earlier
superseded result books.  It copies only the final experiment result workbooks
and creates one long CSV suitable for statistical analysis.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import shutil
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
DEFAULT_SOURCE = Path(
    "/Users/muminemresenol/Desktop/Osman_Outputs/research_projects/"
    "reformer_pilates_revision_1/data"
)

# Each entry contains a final experimental result source, not an input-instance
# workbook.  Rows are read from its stated run-summary sheet.
SOLUTION_SOURCES = (
    ("reformer_expanded_benchmark_detailed_exports_60s.xlsx", "master_summary", "90 expanded instances; CP-Interval and MIP detailed solutions at 60 s"),
    ("factor_separated_benchmark_results_60s_w1_v2_summary.xlsx", "raw_results", "135 factor-separated instances; CP and MIP at 60 s"),
    ("time_limit_benchmark_w1_summary.xlsx", "raw_results", "90 expanded instances; CP and MIP time-limit study at 60/300/600 s"),
    ("time_limit_benchmark_extra_large_60s_w1.xlsx", "raw_results", "30 extra-large instances; CP and MIP at 60 s"),
    ("time_limit_benchmark_extra_large_300_600s_w1.xlsx", "raw_results", "30 extra-large instances; CP and MIP at 300/600 s"),
    ("cp_repeats_original_300_600_w1.xlsx", "raw_runs", "CP repeated runs: 90 expanded instances, 5 repeats at 300/600 s"),
    ("cp_repeats_extra_large_60_300_600_w1.xlsx", "raw_runs", "CP repeated runs: 30 extra-large instances, 5 repeats at 60/300/600 s"),
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-dir", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output-dir", type=Path, default=ROOT)
    args = parser.parse_args()
    archive = args.output_dir / "workbooks"
    archive.mkdir(parents=True, exist_ok=True)
    runs, catalog = [], []
    for filename, sheet_name, scope in SOLUTION_SOURCES:
        source = args.source_dir / filename
        if not source.exists():
            raise FileNotFoundError(source)
        destination = archive / filename
        shutil.copy2(source, destination)
        workbook = load_workbook(source, read_only=True, data_only=True)
        sheet = workbook[sheet_name]
        values = sheet.iter_rows(values_only=True)
        headers = next(values)
        headers = [str(h) if h is not None else f"unnamed_{i}" for i, h in enumerate(headers)]
        row_count = 0
        for row in values:
            if not any(value is not None for value in row):
                continue
            record = {headers[i]: row[i] for i in range(min(len(headers), len(row)))}
            runs.append({"source_file": filename, "source_sheet": sheet_name, **record})
            row_count += 1
        catalog.append({
            "source_file": filename, "source_sheet": sheet_name, "run_rows": row_count,
            "scope": scope, "sha256": sha256(source),
        })

    fields = sorted({key for row in runs for key in row})
    with (args.output_dir / "solution_runs_long.csv").open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="ignore")
        writer.writeheader(); writer.writerows(runs)
    with (args.output_dir / "solution_catalog.csv").open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(catalog[0]))
        writer.writeheader(); writer.writerows(catalog)
    print(f"archived workbooks: {len(catalog)}")
    print(f"solution-run rows: {len(runs)}")


if __name__ == "__main__":
    main()
