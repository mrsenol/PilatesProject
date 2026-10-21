#!/usr/bin/env python3
"""Build the canonical Reformer Pilates scheduling benchmark dataset.

The input workbooks are kept outside this directory because they are the
authoritative, sheet-oriented experiment inputs.  This script normalizes every
instance sheet into analysis-friendly CSV tables without altering any values.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
DEFAULT_SOURCE = Path(
    "/Users/muminemresenol/Desktop/Osman_Outputs/research_projects/"
    "reformer_pilates_revision_1/data"
)
SOURCE_BOOKS = (
    "reformer_benchmark_instances_expanded.xlsx",
    "reformer_benchmark_instances_extra_large.xlsx",
    "reformer_factor_separated_benchmark.xlsx",
)


def write_csv(path: Path, rows: list[dict], fields: list[str]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def source_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def worksheet_sections(ws):
    """Yield (section, header, records) from the vertically structured sheet."""
    rows = list(ws.iter_rows(values_only=True))
    index = 0
    while index < len(rows):
        row = rows[index]
        if row[0] != "section":
            index += 1
            continue
        section = row[1]
        index += 1
        if section == "meta":
            values = {}
            while index < len(rows) and rows[index][0] is not None:
                values[rows[index][0]] = rows[index][1]
                index += 1
            yield section, values, []
            continue
        header = [value for value in rows[index] if value is not None]
        index += 1
        records = []
        while index < len(rows) and rows[index][0] is not None:
            records.append(dict(zip(header, rows[index][: len(header)])))
            index += 1
        yield section, header, records


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-dir", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output-dir", type=Path, default=ROOT / "canonical")
    args = parser.parse_args()
    output = args.output_dir
    output.mkdir(parents=True, exist_ok=True)

    instances, instructors, requests, pairs, patterns, sources = [], [], [], [], [], []
    for book_name in SOURCE_BOOKS:
        book = args.source_dir / book_name
        if not book.exists():
            raise FileNotFoundError(book)
        workbook = load_workbook(book, read_only=True, data_only=True)
        sources.append({"source_book": book_name, "sha256": source_sha256(book)})
        for ws in workbook.worksheets:
            if ws.title == "summary":
                continue
            meta, sections = {}, {}
            for section, header, records in worksheet_sections(ws):
                if section == "meta":
                    meta = header
                else:
                    sections[section] = records
            instance_id = meta["instance_id"]
            family_parts = instance_id.split("_")
            # Factor-separated IDs begin s/m/l; expanded IDs use words.
            family = {"s": "small", "m": "medium", "l": "large"}.get(
                family_parts[0], family_parts[0]
            )
            instances.append({
                "instance_id": instance_id, "source_book": book_name,
                "seed": meta["seed"], "machines": meta["machines"],
                "max_slots_per_instructor_per_day": meta["max_slots_per_instructor_per_day"],
                "family": family,
            })
            for record in sections.get("instructors", []):
                instructors.append({"instance_id": instance_id, **record})
            for section, request_kind in (("individual_requests", "individual"), ("group_customers", "group")):
                for record in sections.get(section, []):
                    requests.append({"instance_id": instance_id, "request_kind": request_kind, **record})
            for record in sections.get("candidate_pairs", []):
                pairs.append({"instance_id": instance_id, **record})
            for record in sections.get("day_patterns", []):
                patterns.append({"instance_id": instance_id, **record})

    write_csv(output / "instances.csv", instances, list(instances[0]))
    write_csv(output / "instructors.csv", instructors, list(instructors[0]))
    write_csv(output / "requests.csv", requests, list(requests[0]))
    write_csv(output / "candidate_pairs.csv", pairs, list(pairs[0]))
    write_csv(output / "day_patterns.csv", patterns, list(patterns[0]))
    write_csv(output / "sources.csv", sources, ["source_book", "sha256"])
    print("\n".join(f"{name}: {len(rows)} rows" for name, rows in {
        "instances": instances, "instructors": instructors, "requests": requests,
        "candidate_pairs": pairs, "day_patterns": patterns,
    }.items()))


if __name__ == "__main__":
    main()
